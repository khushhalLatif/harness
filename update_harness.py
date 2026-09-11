"""
SCRIPT 3 of 3 — pushes approved ssc_appname/ssc_appversion values to Harness.

Reads harness_gitlab_master.xlsx. For every row:

  - If "Approved for Update (Y/N)" is exactly Y (case-insensitive) -> pushes
    Appname_final and Appversion_final to that service's Advanced >
    Variables on Harness (ssc_appname, ssc_appversion) — creating the
    variable if it doesn't exist yet, updating it if it does.
  - Anything else (N, blank, or anything other than Y) -> skipped entirely,
    no Harness API call made for that row at all.

SAFETY — DRY_RUN defaults to True. Every approved row is fully evaluated
and the exact change that WOULD be made is written to Comments, but no
PUT is actually sent to Harness. Review a dry run first; only flip
DRY_RUN to False once you're confident in the Approved/Final values.

This script never touches "Done by" — that's yours to fill in by hand.
Its own result for each row is written into the (single) Comments column,
using the same tagged-note pattern Script 1 and Script 2 already use, so
nothing from either of those is overwritten.

Run Script 1 and Script 2 first, fill in Approved for Update / Appname_final
/ Appversion_final by hand, THEN run this.

pip install requests pyyaml openpyxl --break-system-packages
"""

import os
import sys
import time
import logging
from datetime import datetime

import yaml
import requests
from openpyxl import load_workbook

from harness_config import load_config, require_harness

# ===========================================================================
# CONFIG
# ===========================================================================

_config = load_config()
_harness = require_harness(_config)

HARNESS_BASE_URL = _harness["base_url"]
HARNESS_API_KEY = _harness["api_key"]
HARNESS_ACCOUNT_ID = _harness["account_id"]
HARNESS_ORG_ID = _harness["org_id"]
HARNESS_PROJECT_ID = _harness["project_id"]

MASTER_EXCEL_FILE = "harness_gitlab_master.xlsx"
LOG_FILE = f"script3_harness_update_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

DRY_RUN = True  # False = actually push updates to Harness. Read the module docstring first.
SLEEP_BETWEEN_ROWS = 0.3
MAX_RETRIES = 3
RETRY_BACKOFF_SECONDS = 2

# ===========================================================================
# Shared column schema — MUST exactly match Script 1 / Script 2's ALL_COLUMNS.
# ===========================================================================

ALL_COLUMNS = [
    "Service Identifier", "Service Name",
    "Appversion_final", "Appname_final",
    "Comments",
    "Service Url",
    "Existing_Branch",
    "Corresponding Gitlab component1", "GitLab File Path1",
    "Corresponding Gitlab component2", "GitLab File Path2",
    "New_ssc_appversion1", "New_ssc_appversion2", "Existing_ssc_appversion",
    "Done by",
    "Artifact Path",
    "New_ssc_appname1", "New_ssc_appname2", "Existing_ssc_appname",
    "Approved for Update (Y/N)",
]

# ===========================================================================
# Logging
# ===========================================================================

logger = logging.getLogger("script3")
logger.setLevel(logging.DEBUG)
fh = logging.FileHandler(LOG_FILE, encoding="utf-8")
fh.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
ch = logging.StreamHandler()
ch.setLevel(logging.INFO)
ch.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
logger.addHandler(fh)
logger.addHandler(ch)

# ===========================================================================
# Harness API helpers
# ===========================================================================

HEADERS = {
    "x-api-key": HARNESS_API_KEY,
    "Content-Type": "application/json",
    # Same corporate-network-friendly identity used in Script 1 — some
    # SSL-inspecting proxies terminate connections carrying the default
    # python-requests identifier.
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"),
}
COMMON_PARAMS = {
    "accountIdentifier": HARNESS_ACCOUNT_ID,
    "orgIdentifier": HARNESS_ORG_ID,
    "projectIdentifier": HARNESS_PROJECT_ID,
}


def request_with_retry(method: str, url: str, **kwargs) -> requests.Response:
    last_exc = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.request(method, url, timeout=30, **kwargs)
            if resp.status_code >= 500:
                raise requests.HTTPError(f"{resp.status_code} server error")
            return resp
        except requests.RequestException as exc:
            last_exc = exc
            logger.warning("Request failed (attempt %s/%s) %s %s: %s",
                            attempt, MAX_RETRIES, method, url, exc)
            time.sleep(RETRY_BACKOFF_SECONDS * attempt)
    raise last_exc


def get_service_yaml(service_identifier: str) -> dict:
    url = f"{HARNESS_BASE_URL}/ng/api/servicesV2/{service_identifier}"
    resp = request_with_retry("GET", url, headers=HEADERS, params=COMMON_PARAMS)
    resp.raise_for_status()
    raw_yaml = resp.json()["data"]["service"]["yaml"]
    return yaml.safe_load(raw_yaml)


def put_service_yaml(service_identifier: str, yaml_dict: dict) -> None:
    new_yaml_str = yaml.dump(yaml_dict, sort_keys=False)
    body = {
        "identifier": service_identifier,
        "orgIdentifier": HARNESS_ORG_ID,
        "projectIdentifier": HARNESS_PROJECT_ID,
        "yaml": new_yaml_str,
    }
    resp = request_with_retry("PUT", f"{HARNESS_BASE_URL}/ng/api/servicesV2",
                               headers=HEADERS, params=COMMON_PARAMS, json=body)
    resp.raise_for_status()


def upsert_variable(yaml_dict: dict, var_name: str, new_value: str) -> tuple[str, object]:
    """Returns (action, old_value) where action is 'created' or 'updated'."""
    variables = yaml_dict["service"]["serviceDefinition"]["spec"].setdefault("variables", [])
    for var in variables:
        if var.get("name") == var_name:
            old_value = var.get("value")
            var["value"] = new_value
            return "updated", old_value
    variables.append({"name": var_name, "type": "String", "value": new_value})
    return "created", None


# ===========================================================================
# Master Excel file helpers
# ===========================================================================


def get_cell(ws, row_idx: int, column_name: str):
    col_idx = ALL_COLUMNS.index(column_name) + 1
    return ws.cell(row=row_idx, column=col_idx).value


def set_script_comment(ws, row_idx: int, script_tag: str, text: str):
    """Replaces this script's previous comment(s) on the row rather than
    stacking new ones on top; Script 1/2's (and your own) notes are left
    completely untouched."""
    col_idx = ALL_COLUMNS.index("Comments") + 1
    cell = ws.cell(row=row_idx, column=col_idx)
    existing_lines = (cell.value or "").split("\n")
    other_lines = [line for line in existing_lines if line and not line.startswith(f"[{script_tag} ")]
    if text:
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        other_lines.append(f"[{script_tag} {stamp}] {text}")
    cell.value = "\n".join(other_lines) if other_lines else None


# ===========================================================================
# Per-row processing
# ===========================================================================


def process_row(ws, row_idx: int) -> str:
    svc_id = get_cell(ws, row_idx, "Service Identifier")
    svc_name = get_cell(ws, row_idx, "Service Name") or svc_id or "(unknown)"
    approved_raw = get_cell(ws, row_idx, "Approved for Update (Y/N)")
    approved = (approved_raw or "").strip().upper()

    if approved != "Y":
        # No comment written for the common case (not approved yet) — keeps
        # the Comments column from filling up with "skipped" on every row
        # of a 330-row sheet where most rows simply aren't approved yet.
        return "skipped"

    if not svc_id:
        set_script_comment(ws, row_idx, "Script3",
                            "ERROR: Approved=Y but Service Identifier is blank")
        return "error"

    appname_final = get_cell(ws, row_idx, "Appname_final")
    appversion_final = get_cell(ws, row_idx, "Appversion_final")

    if not appname_final or not appversion_final:
        set_script_comment(ws, row_idx, "Script3",
                            "ERROR: Approved=Y but Appname_final/Appversion_final is blank "
                            "— refusing to push an empty value")
        return "error"

    try:
        yaml_dict = get_service_yaml(svc_id)
        action_name, old_name = upsert_variable(yaml_dict, "ssc_appname", str(appname_final))
        action_version, old_version = upsert_variable(yaml_dict, "ssc_appversion", str(appversion_final))

        summary = (f"ssc_appname {old_name!r}->{appname_final!r} ({action_name}); "
                   f"ssc_appversion {old_version!r}->{appversion_final!r} ({action_version})")

        if DRY_RUN:
            set_script_comment(ws, row_idx, "Script3", f"DRY RUN — would update: {summary}")
            return "dry_run"

        put_service_yaml(svc_id, yaml_dict)
        set_script_comment(ws, row_idx, "Script3", f"Updated on Harness: {summary}")
        return "updated"

    except Exception as exc:  # noqa: BLE001 — keep going across all rows
        logger.error("Failed on service '%s': %s", svc_name, exc)
        set_script_comment(ws, row_idx, "Script3", f"ERROR: {exc}")
        return "error"


# ===========================================================================
# Main
# ===========================================================================


def main():
    mode = "DRY RUN (no changes will be pushed to Harness)" if DRY_RUN else "LIVE — pushing real updates to Harness"
    logger.info("=== Script 3 (Harness update) started — %s ===", mode)

    if not os.path.exists(MASTER_EXCEL_FILE):
        sys.exit(f"{MASTER_EXCEL_FILE} not found — run Script 1 and Script 2 first.")

    wb = load_workbook(MASTER_EXCEL_FILE)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if header != ALL_COLUMNS:
        sys.exit(
            f"{MASTER_EXCEL_FILE}'s header doesn't match the expected schema.\n"
            f"Expected: {ALL_COLUMNS}\nFound:    {header}"
        )

    total_rows = ws.max_row - 1
    counters = {}

    for row_idx in range(2, ws.max_row + 1):
        svc_name = get_cell(ws, row_idx, "Service Name") or "(unknown)"
        try:
            result = process_row(ws, row_idx)
        except Exception as exc:  # noqa: BLE001
            logger.error("Unhandled error on row %s: %s", row_idx, exc)
            set_script_comment(ws, row_idx, "Script3", f"ERROR: {exc}")
            result = "error"
        counters[result] = counters.get(result, 0) + 1

        if result != "skipped":
            logger.info("[row %s/%s] %s -> %s", row_idx - 1, total_rows, svc_name, result)

        time.sleep(SLEEP_BETWEEN_ROWS)

    wb.save(MASTER_EXCEL_FILE)
    logger.info("=== Script 3 complete: %s ===", counters)
    print(f"\nDone ({mode}).")
    print(f"Summary: {counters}")
    print(f"Master file: {MASTER_EXCEL_FILE}")
    print(f"Log file:    {LOG_FILE}")
    if DRY_RUN and counters.get("dry_run"):
        print(f"\n{counters['dry_run']} row(s) would be updated. Review the Comments column,")
        print("then set DRY_RUN = False at the top of this script to actually push them.")


if __name__ == "__main__":
    main()
